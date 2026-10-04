#!/usr/bin/env python3
"""Turn a Canva export (PDF, or a .zip of page images) into a .pptx,
one full-slide image per page.

Usage:
    python3 diakonverter.py DECK.pdf [-o OUTPUT.pptx] [--width 3840] [--format jpeg|png]
    python3 diakonverter.py DECK.zip [-o OUTPUT.pptx]

The output defaults to "<source name> converted.pptx" next to the source.

PDF pages are rendered at --width pixels wide. Zip images are used as-is, in
the order they are stored in the zip: Canva names titled pages by their title
rather than their number, so sorting by filename would scramble them.
The slide's aspect ratio matches the first page.

Requires: pip install python-pptx pymupdf
"""
import argparse
import io
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
SLIDE_WIDTH = Inches(13.333)  # PowerPoint's standard widescreen width


def images_from_zip(zip_path):
    with zipfile.ZipFile(zip_path) as archive:
        for entry in archive.infolist():
            name = Path(entry.filename)
            if (entry.is_dir() or name.suffix.lower() not in IMAGE_EXTENSIONS
                    or name.name.startswith(".") or "__MACOSX" in name.parts):
                continue
            yield io.BytesIO(archive.read(entry))


def images_from_pdf(pdf_path, width, fmt, quality):
    import pymupdf

    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            zoom = width / page.rect.width
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
            if fmt == "png":
                yield io.BytesIO(pix.tobytes("png"))
            else:
                yield io.BytesIO(pix.tobytes("jpeg", jpg_quality=quality))


def add_image_slide(prs, image):
    with Image.open(image) as img:
        img_w, img_h = img.size
    image.seek(0)

    if len(prs.slides) == 0:
        prs.slide_width = SLIDE_WIDTH
        prs.slide_height = Emu(int(SLIDE_WIDTH * img_h / img_w))

    # Scale to fit inside the slide, preserving aspect ratio, then center.
    scale = min(prs.slide_width / img_w, prs.slide_height / img_h)
    width, height = int(img_w * scale), int(img_h * scale)
    left = (prs.slide_width - width) // 2
    top = (prs.slide_height - height) // 2
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    slide.shapes.add_picture(image, left, top, width, height)


def main():
    parser = argparse.ArgumentParser(description="Convert a Canva PDF or image zip into a .pptx")
    parser.add_argument("source", type=Path, help="a Canva PDF export, or a .zip of PNG/JPEG pages")
    parser.add_argument("-o", "--output", type=Path,
                        help='output file (default: "<source name> converted.pptx" next to the source)')
    parser.add_argument("--width", type=int, default=3840,
                        help="PDF only: rendered image width in pixels (default: 3840)")
    parser.add_argument("--format", choices=["jpeg", "png"], default="jpeg",
                        help="PDF only: image format; png is lossless but ~7x larger (default: jpeg)")
    parser.add_argument("--quality", type=int, default=95,
                        help="PDF only: JPEG quality 1-100 (default: 95)")
    args = parser.parse_args()

    source = args.source
    kind = source.suffix.lower()
    if not source.is_file() or kind not in {".pdf", ".zip"}:
        sys.exit(f"Expected a .pdf or .zip file: {source}")
    if kind == ".pdf":
        images = images_from_pdf(source, args.width, args.format, args.quality)
    else:
        images = images_from_zip(source)

    prs = Presentation()
    for number, image in enumerate(images, 1):
        add_image_slide(prs, image)
        print(f"\rSlide {number}", end="", flush=True)
    print()
    if len(prs.slides) == 0:
        sys.exit(f"No pages or images found in {source}")

    output = args.output or source.with_name(f"{source.stem} converted.pptx")
    prs.save(output)
    print(f"Created {output} with {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
